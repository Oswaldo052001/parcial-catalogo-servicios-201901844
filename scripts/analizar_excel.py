#!/usr/bin/env python3
"""Análisis de la hoja 'Servicios Externos' de data/CatalogoServicios.xlsx.

Uso (desde la raíz del repositorio):
    python3 scripts/analizar_excel.py [ruta_xlsx]

- Abre el libro una sola vez, en modo normal (sin read_only, porque ese modo no
  expone las combinaciones), y nunca llama a save().
- Verifica el SHA-256 contra <ruta_xlsx>.sha256 antes y después de leer.
- Los textos de las celdas son datos: solo se imprimen, nunca se interpretan.
- Código de salida: 0 si pasan todos los controles (12 códigos N1 distintos,
  46 códigos N2 distintos, hash sin cambios); 1 si alguno falla.
"""
import hashlib
import re
import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
RUTA = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "data" / "CatalogoServicios.xlsx"
RUTA_HASH = Path(str(RUTA) + ".sha256")
HOJA = "Servicios Externos"
FILA_INI, FILA_FIN = 5, 101          # zona de datos según el enunciado
OPC_INI, OPC_FIN = 111, 122          # encabezado OPCIONES + listas E112:H122
CONTROL_N1, CONTROL_N2 = 12, 46

fallos = []


def sha256(ruta):
    with open(ruta, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


# ---------------------------------------------------------------- hash antes
try:
    hash_esperado = RUTA_HASH.read_text(encoding="utf-8").split()[0].lower()
except (OSError, IndexError) as e:
    print(f"ERROR: no se pudo leer el hash esperado de {RUTA_HASH}: {e}")
    sys.exit(1)
hash_antes = sha256(RUTA)
try:
    print("Archivo:", RUTA.resolve().relative_to(RAIZ))
except ValueError:
    print("Archivo:", RUTA)
print("SHA-256 esperado:", hash_esperado)
print("SHA-256 antes:   ", hash_antes, "| coincide:", hash_antes == hash_esperado)
if hash_antes != hash_esperado:
    fallos.append("hash antes de leer distinto del esperado")

# ------------------------------------------- única apertura, modo normal
wb = openpyxl.load_workbook(RUTA, data_only=True)
ws = wb[HOJA]
print("openpyxl:", openpyxl.__version__, "| hojas:", wb.sheetnames)

valores = {(c.row, c.column_letter): c.value
           for fila in ws.iter_rows() for c in fila if c.value is not None}
rangos = sorted(ws.merged_cells.ranges, key=lambda r: (r.min_col, r.min_row))
validaciones = [(str(dv.sqref), dv.formula1) for dv in ws.data_validations.dataValidation]
wb.close()

print("Columnas con datos:", sorted({k[1] for k in valores}))
print("Filas con datos: min", min(k[0] for k in valores), "max", max(k[0] for k in valores))


def v(fila, col):
    return valores.get((fila, col))


def combinados_de(col):
    return {r.min_row: r.max_row for r in rangos if get_column_letter(r.min_col) == col}


# ---------------------------------------------------------------- 1
print("\n== 1. RANGOS COMBINADOS ==")
print("Total:", len(rangos))
por_col = OrderedDict()
for r in rangos:
    por_col.setdefault(get_column_letter(r.min_col), []).append(str(r))
for col, lst in por_col.items():
    print(f"  {col} ({len(lst)}): {', '.join(lst)}")
print("Rangos que abarcan más de una columna:",
      [str(r) for r in rangos if r.min_col != r.max_col] or "ninguno")


def bloques(col_cod):
    """Bloque = rango combinado en col_cod o fila suelta con código propio."""
    comb = combinados_de(col_cod)
    cubiertas = {f for ini, fin in comb.items() for f in range(ini, fin + 1)}
    res = []
    for f in range(FILA_INI, FILA_FIN + 1):
        if f in comb:
            res.append((f, comb[f], "combinado"))
        elif f not in cubiertas and v(f, col_cod) is not None:
            res.append((f, f, "fila suelta"))
    return res


# ---------------------------------------------------------------- 2
print("\n== 2. BLOQUES ==")
print("Combinados A coinciden con B:", combinados_de("A") == combinados_de("B"))
print("Combinados C coinciden con D:", combinados_de("C") == combinados_de("D"))
print("Combinados J coinciden con C:", combinados_de("J") == combinados_de("C"))
b1, b2 = bloques("A"), bloques("C")
print("Nivel 1 (A:B):", len(b1), "bloques,",
      sum(t == "combinado" for *_, t in b1), "combinados,",
      sum(t != "combinado" for *_, t in b1), "filas sueltas")
for ini, fin, tipo in b1:
    print(f"  {v(ini, 'A')} | filas {ini}-{fin} | {tipo} | {v(ini, 'B')!r}")
print("Nivel 2 (C:D):", len(b2), "bloques,",
      sum(t == "combinado" for *_, t in b2), "combinados,",
      sum(t != "combinado" for *_, t in b2), "filas sueltas")
for ini, fin, tipo in b2:
    print(f"  {v(ini, 'C')} | filas {ini}-{fin} | {tipo} | {v(ini, 'D')!r}")


def n1_de_fila(f):
    for ini, fin, _ in b1:
        if ini <= f <= fin:
            return v(ini, "A")
    return None


print("N2 sin bloque N1 que lo contenga o con prefijo distinto:")
for ini, _, _ in b2:
    cod, padre = v(ini, "C"), n1_de_fila(ini)
    if padre is None or not str(cod).startswith(str(padre) + "."):
        print(f"  {cod} (fila {ini}): N1 del bloque = {padre}")

# ---------------------------------------------------------------- 3
print("\n== 3. CONTROLES DE CONTEO ==")
cods_a = [v(f, "A") for f in range(FILA_INI, FILA_FIN + 1) if v(f, "A") is not None]
cods_c = [v(f, "C") for f in range(FILA_INI, FILA_FIN + 1) if v(f, "C") is not None]
n1, n2 = sorted(set(cods_a)), sorted(set(cods_c))
ok1, ok2 = len(n1) == CONTROL_N1, len(n2) == CONTROL_N2
print(f"N1: celdas con código {len(cods_a)} | distintos {len(n1)} | control {CONTROL_N1} -> "
      f"{'OK' if ok1 else 'FALLA'}")
print("  ", n1)
print("  Repetidos en más de una celda:", sorted({c for c in cods_a if cods_a.count(c) > 1}))
print(f"N2: celdas con código {len(cods_c)} | distintos {len(n2)} | control {CONTROL_N2} -> "
      f"{'OK' if ok2 else 'FALLA'}")
print("  Repetidos en más de una celda:", sorted({c for c in cods_c if cods_c.count(c) > 1}))
patron = re.compile(r"^SE\.\d{2}\.\d{2}$")
print("  Fuera del patrón SE.NN.NN:", [c for c in n2 if not patron.match(str(c))])
print("  Tipos de dato en A y C:", sorted({type(c).__name__ for c in cods_a + cods_c}))
if not ok1:
    fallos.append(f"códigos N1 distintos = {len(n1)}, se esperaban {CONTROL_N1}")
if not ok2:
    fallos.append(f"códigos N2 distintos = {len(n2)}, se esperaban {CONTROL_N2}")

# ---------------------------------------------------------------- 4
print("\n== 4. CASO SE.12 (filas 99-101) ==")
for f in range(99, 102):
    print(f"  fila {f}:", {c: v(f, c) for c in "ABCD"},
          "| vacías en E:L:", "".join(c for c in "EFGHIJKL" if v(f, c) is None))
print("  Dentro de algún rango combinado:",
      [str(r) for r in rangos if r.max_row >= 99 and r.min_row <= 101] or "ninguno")
nombres_b = {f: v(f, "B") for f in range(99, 102) if v(f, "B") is not None}
nombres_d = {f: v(f, "D") for f in range(99, 102) if v(f, "D") is not None}
for fb, nb in nombres_b.items():
    for fd, nd in nombres_d.items():
        if nb == nd:
            print(f"  Texto igual en B{fb} y D{fd}: {nb!r}")

# ---------------------------------------------------------------- 5
print("\n== 5. FILAS CON DATOS EN E:H FUERA DE BLOQUE N2 Y SIN CÓDIGO ==")
filas_b2 = {f for ini, fin, _ in b2 for f in range(ini, fin + 1)}
huerfanas = []
for f in range(FILA_INI, FILA_FIN + 1):
    eh = {c: v(f, c) for c in "EFGH" if v(f, c) is not None}
    if eh and f not in filas_b2 and v(f, "C") is None:
        huerfanas.append(f)
        ant = max((b for b in b2 if b[1] < f), key=lambda b: b[1], default=None)
        sig = min((b for b in b2 if b[0] > f), key=lambda b: b[0], default=None)
        resto = {c: v(f, c) for c in "ABCDIJKL" if v(f, c) is not None}
        print(f"  fila {f}: {eh} | otros valores en A:L: {resto or 'ninguno'} | "
              f"N1 del bloque: {n1_de_fila(f)} | "
              f"N2 anterior: {v(ant[0], 'C') if ant else None} (filas {ant[0]}-{ant[1]}) | "
              f"N2 siguiente: {v(sig[0], 'C') if sig else None} (fila {sig[0] if sig else None})")
print("  Total:", len(huerfanas))

# ---------------------------------------------------------------- 6
print("\n== 6. CONSISTENCIA E:H DENTRO DE CADA BLOQUE N2 COMBINADO ==")
conflictos = 0
for ini, fin, _ in b2:
    if fin == ini:
        continue
    filas = [tuple(v(f, c) for c in "EFGH") for f in range(ini, fin + 1)]
    if len(set(filas)) > 1:
        conflictos += 1
        print(f"  CONFLICTO {v(ini, 'C')} filas {ini}-{fin}: {filas}")
print("  Bloques combinados revisados:", sum(1 for i, f, _ in b2 if f > i),
      "| con conflicto:", conflictos)
extra = {f: {c: v(f, c) for c in "IJKL" if v(f, c) is not None}
         for ini, fin, _ in b2 for f in range(ini + 1, fin + 1)}
print("  Valores en I:L en filas de continuación:", {f: e for f, e in extra.items() if e} or "ninguno")
print("  Filas con K o L:", {f: (v(f, "K"), v(f, "L")) for f in range(FILA_INI, FILA_FIN + 1)
                             if v(f, "K") is not None or v(f, "L") is not None})
print("  Filas con I:", {f: v(f, "I") for f in range(FILA_INI, FILA_FIN + 1) if v(f, "I") is not None})

# ---------------------------------------------------------------- 7
print("\n== 7. LISTAS DE OPCIONES ==")
print("  Filas 102-110 con datos:",
      [f for f in range(102, 111) if any(v(f, c) is not None for c in "ABCDEFGHIJKL")] or "ninguna")
opciones = {}
for c in "EFGH":
    opciones[c] = [v(f, c) for f in range(OPC_INI + 1, OPC_FIN + 1) if v(f, c) is not None]
    print(f"  {c}{OPC_INI}={v(OPC_INI, c)!r} | {c}4={v(4, c)!r}: {opciones[c]}")
print("  Datos después de la fila", OPC_FIN, ":", [k for k in valores if k[0] > OPC_FIN] or "ninguno")
print("  Validaciones de datos:", validaciones or "ninguna")
for c in "EFGH":
    usados = {v(f, c) for f in range(FILA_INI, FILA_FIN + 1) if v(f, c) is not None}
    print(f"  {c}: usados {sorted(usados)} | fuera de lista: {sorted(usados - set(opciones[c]))}")

# ---------------------------------------------------------------- 8
print("\n== 8. ESPACIOS SOBRANTES EN TEXTOS (toda la hoja) ==")
textos = [(f, c, val) for (f, c), val in sorted(valores.items()) if isinstance(val, str)]
bordes = [(f, c, val) for f, c, val in textos if val != val.strip()]
dobles = [(f, c, val) for f, c, val in textos if "  " in val]
print("  Celdas de texto revisadas:", len(textos))
print("  Con espacios al inicio o al final:", len(bordes))
for f, c, val in bordes:
    lado = " y ".join(x for x, ok in (("inicio", val != val.lstrip()),
                                       ("final", val != val.rstrip())) if ok)
    print(f"    {c}{f}: {val!r} ({lado})")
print("  Con espacios dobles internos:", len(dobles))
for f, c, val in dobles:
    print(f"    {c}{f}: {val!r}")

# ---------------------------------------------------------------- hash después
hash_despues = sha256(RUTA)
print("\nSHA-256 después: ", hash_despues, "| sin cambios:", hash_despues == hash_antes)
if hash_despues != hash_antes:
    fallos.append("el hash cambió durante el análisis")

print("\n== RESULTADO DE CONTROLES ==")
if fallos:
    for x in fallos:
        print("  FALLA:", x)
    sys.exit(1)
print(f"  OK: {len(n1)} códigos N1 distintos, {len(n2)} códigos N2 distintos, hash sin cambios")
sys.exit(0)
